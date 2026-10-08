"""Unit tests: the sheet-set library (xlsx writer/reader, lint rules, build record) and the generated mach set.

Run: python3 -m unittest discover -s tests (python 3.9+). Scratch folders live in <repo>/.tmp/ (never /tmp).
openpyxl is optional: when it is importable the workbooks are also loaded with it, otherwise that test is skipped.
"""
import os
import re
import shutil
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "lib"))
sys.dont_write_bytecode = True

import mach_sheets as sh  # noqa: E402
import mach_sheets_gen as gen  # noqa: E402

NS = {"m": sh.NS_MAIN}
SET_DIR = os.path.join(ROOT, "docs", "sheets")


def scratch():
    os.makedirs(os.path.join(ROOT, ".tmp"), exist_ok=True)
    return tempfile.mkdtemp(prefix="sheets.", dir=os.path.join(ROOT, ".tmp"))


def make_set(d, files, names=None):
    """Write CSV files {name: (header, rows)} and an index.csv for them; return the folder."""
    os.makedirs(d, exist_ok=True)
    idx = []
    for i, (fn, (header, rows)) in enumerate(sorted(files.items())):
        with open(os.path.join(d, fn), "w", encoding="utf-8", newline="") as fh:
            fh.write(sh.write_csv_text(header, rows))
        sheet = (names or {}).get(fn, fn[:-4].capitalize())
        idx.append(["s%d" % i, fn, sheet, fn, "", "me", "", "", "", ""])
    with open(os.path.join(d, "index.csv"), "w", encoding="utf-8", newline="") as fh:
        fh.write(sh.write_csv_text(sh.INDEX_COLS, idx))
    return d


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = scratch()
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def build(self, d, today="2026-10-08"):
        problems, plan = sh.plan_build(d, today)
        self.assertEqual(problems, [])
        with open(os.path.join(d, "index.csv"), "w", encoding="utf-8", newline="") as fh:
            fh.write(plan["index_text"])
        with open(plan["xlsx_path"], "wb") as fh:
            fh.write(plan["xlsx_bytes"])
        return plan

    def levels(self, d):
        return sorted(set(lv for lv, _ in sh.lint_set(d)))


class XlsxStructure(Base):
    def setUp(self):
        Base.setUp(self)
        self.d = make_set(os.path.join(self.tmp, "proj"), {
            "tasks.csv": (["id", "name", "note"], [["1", "a & b", "x < y"], ["2", "  lead space", "line1\nline2"], ["3", "café ✓", "=SUM(A1)"]]),
            "people.csv": (["who", "role"], [["ann", "dev"]]),
        }, {"tasks.csv": "Task list", "people.csv": "People"})
        self.plan = self.build(self.d)
        self.xlsx = os.path.join(self.d, "proj.xlsx")

    def test_workbook_name_defaults_to_folder_name(self):
        self.assertEqual(sh.workbook_path(self.d), self.xlsx)
        self.assertTrue(os.path.isfile(self.xlsx))

    def test_sheet_order_index_first_then_csv_order(self):
        self.assertEqual([n for n, _ in sh.read_xlsx(self.xlsx)], ["Index", "People", "Task list"])

    def test_zip_parts_are_wellformed_xml(self):
        with zipfile.ZipFile(self.xlsx) as z:
            self.assertIsNone(z.testzip())
            names = z.namelist()
            self.assertEqual(names[0], "[Content_Types].xml")
            for n in names:
                ET.fromstring(z.read(n))
            for n in ("xl/workbook.xml", "xl/styles.xml", "xl/worksheets/sheet1.xml", "docProps/custom.xml"):
                self.assertIn(n, names)

    def test_header_row_is_frozen_filtered_and_bold(self):
        with zipfile.ZipFile(self.xlsx) as z:
            for i in (1, 2, 3):
                sheet = ET.fromstring(z.read("xl/worksheets/sheet%d.xml" % i))
                pane = sheet.find("m:sheetViews/m:sheetView/m:pane", NS)
                self.assertEqual((pane.get("ySplit"), pane.get("topLeftCell"), pane.get("state")), ("1", "A2", "frozen"))
                flt = sheet.find("m:autoFilter", NS)
                self.assertTrue(re.match(r"^A1:[A-Z]+\d+$", flt.get("ref")))
                first = sheet.find("m:sheetData/m:row", NS)
                self.assertTrue(all(c.get("s") == "1" for c in first))
                self.assertTrue(len(sheet.findall("m:cols/m:col", NS)) >= 2)
            styles = ET.fromstring(z.read("xl/styles.xml"))
            fonts = styles.findall("m:fonts/m:font", NS)
            xfs = styles.findall("m:cellXfs/m:xf", NS)
            self.assertIsNotNone(fonts[int(xfs[1].get("fontId"))].find("m:b", NS))

    def test_autofilter_covers_all_rows_and_columns(self):
        with zipfile.ZipFile(self.xlsx) as z:
            tasks = ET.fromstring(z.read("xl/worksheets/sheet3.xml"))
        self.assertEqual(tasks.find("m:autoFilter", NS).get("ref"), "A1:C4")

    def test_cells_round_trip_through_the_reader(self):
        sheets = dict(sh.read_xlsx(self.xlsx))
        self.assertEqual(sheets["Task list"][0], ["id", "name", "note"])
        self.assertEqual(sheets["Task list"][1], ["1", "a & b", "x < y"])
        self.assertEqual(sheets["Task list"][2], ["2", "  lead space", "line1\nline2"])
        self.assertEqual(sheets["Task list"][3], ["3", "café ✓", "=SUM(A1)"])
        # the index sheet mirrors index.csv, with the rows column as a number
        idx = sheets["Index"]
        self.assertEqual(idx[0], sh.INDEX_COLS)
        self.assertEqual([r[2] for r in idx[1:]], ["People", "Task list"])
        self.assertEqual([r[-1] for r in idx[1:]], ["1", "3"])

    def test_formula_like_text_is_stored_as_a_string_never_a_formula(self):
        with zipfile.ZipFile(self.xlsx) as z:
            for n in z.namelist():
                if n.startswith("xl/worksheets/"):
                    self.assertNotIn("<f>", z.read(n).decode("utf-8"))

    def test_same_input_gives_identical_bytes(self):
        again = sh.plan_build(self.d, "2026-12-31")[1]
        with open(self.xlsx, "rb") as fh:
            self.assertEqual(fh.read(), again["xlsx_bytes"])
        self.assertEqual(again["changes"], [])

    def test_zip_has_fixed_timestamps_and_no_clock(self):
        with zipfile.ZipFile(self.xlsx) as z:
            self.assertEqual(set(i.date_time for i in z.infolist()), {(1980, 1, 1, 0, 0, 0)})

    def test_no_absolute_home_paths_in_the_workbook(self):
        with zipfile.ZipFile(self.xlsx) as z:
            blob = b"".join(z.read(n) for n in z.namelist())
        self.assertNotRegex(blob.decode("utf-8"), r"/home/|/Users/")

    @unittest.skipUnless(__import__("importlib").util.find_spec("openpyxl"), "openpyxl not installed")
    def test_openpyxl_can_load_it(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.xlsx)
        self.assertEqual(wb.sheetnames, ["Index", "People", "Task list"])
        ws = wb["Task list"]
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:C4")
        self.assertTrue(ws["A1"].font.b)
        self.assertEqual(ws["B2"].value, "a & b")


class LintRules(Base):
    def good(self):
        return make_set(os.path.join(self.tmp, "set"), {"a.csv": (["x", "y"], [["1", "2"]]), "b.csv": (["k"], [["v"], ["w"]])})

    def messages(self, d):
        return " | ".join(m for _, m in sh.lint_set(d))

    def test_built_set_is_clean(self):
        d = self.good()
        self.build(d)
        self.assertEqual(sh.lint_set(d), [])

    def test_unbuilt_set_warns_only(self):
        d = self.good()
        self.assertEqual(self.levels(d), ["FAIL"])      # rows/updated still empty
        self.build(d)
        os.remove(os.path.join(d, "set.xlsx"))
        self.assertEqual(self.levels(d), ["WARN"])

    def test_orphan_csv_fails(self):
        d = self.good()
        self.build(d)
        shutil.copy(os.path.join(d, "a.csv"), os.path.join(d, "stray.csv"))
        self.assertIn("stray.csv is not listed", self.messages(d))

    def test_missing_listed_file_fails(self):
        d = self.good()
        os.remove(os.path.join(d, "b.csv"))
        self.assertIn("b.csv does not exist", self.messages(d))

    def test_wrong_rows_and_updated_fail(self):
        d = self.good()
        self.build(d)
        with open(os.path.join(d, "a.csv"), "a", encoding="utf-8") as fh:
            fh.write("3,4\n")
        msg = self.messages(d)
        self.assertIn("rows='1' but the CSV has 2", msg)
        with open(os.path.join(d, "index.csv"), encoding="utf-8") as fh:
            text = fh.read().replace("2026-10-08", "yesterday")
        with open(os.path.join(d, "index.csv"), "w", encoding="utf-8") as fh:
            fh.write(text)
        self.assertIn("not a YYYY-MM-DD date", self.messages(d))

    def test_csv_changed_after_build_is_stale(self):
        d = self.good()
        self.build(d)
        with open(os.path.join(d, "b.csv"), "w", encoding="utf-8") as fh:
            fh.write("k\nv\nw\n")           # same bytes shape, same rows -> still current
        self.assertEqual(sh.lint_set(d), [])
        with open(os.path.join(d, "b.csv"), "w", encoding="utf-8") as fh:
            fh.write("k\nv\nz\n")           # same row count, different content
        self.assertIn("changed after set.xlsx was built", self.messages(d))

    def test_sheet_name_rules(self):
        for bad in ("x" * 32, "a:b", "a/b", "a[b", "a*b", "a?b", "'quoted", "History", "index", ""):
            self.assertNotEqual(sh.sheet_name_problem(bad), "", bad)
        for ok in ("x" * 31, "Config keys", "Subsystem units", "a-b_c"):
            self.assertEqual(sh.sheet_name_problem(ok), "", ok)

    def test_duplicate_sheet_names_compare_case_insensitively(self):
        d = make_set(os.path.join(self.tmp, "dup"), {"a.csv": (["x"], []), "b.csv": (["x"], [])}, {"a.csv": "Same", "b.csv": "same"})
        self.assertIn("repeats row", self.messages(d))

    def test_index_columns_must_be_exact(self):
        d = self.good()
        with open(os.path.join(d, "index.csv"), "w", encoding="utf-8") as fh:
            fh.write("id,file,sheet_name\n")
        self.assertIn("columns must be exactly", self.messages(d))

    def test_ragged_and_control_characters_fail(self):
        d = make_set(os.path.join(self.tmp, "bad"), {"a.csv": (["x", "y"], [["1", "2"]])})
        with open(os.path.join(d, "a.csv"), "a", encoding="utf-8") as fh:
            fh.write("only-one-field\n1,2\x01\n")
        msg = self.messages(d)
        self.assertIn("do not have 2 fields", msg)
        self.assertIn("control characters", msg)

    def test_absolute_home_paths_warn(self):
        d = make_set(os.path.join(self.tmp, "p"), {"a.csv": (["path"], [["/home/someone/x"], ["~/ok"], ["/Users/someone/y"]])})
        self.assertIn("absolute home path at A2, A4", self.messages(d))

    def test_formula_like_cells_warn(self):
        d = make_set(os.path.join(self.tmp, "f"), {"a.csv": (["v"], [["=1+1"], ["--apply"], ["-t"], ["@x"]])})
        self.assertIn("2 cell(s) start like a formula", self.messages(d))

    def test_empty_csv_and_empty_header_cell_fail(self):
        d = make_set(os.path.join(self.tmp, "e"), {"a.csv": (["x"], [])})
        open(os.path.join(d, "a.csv"), "w").close()
        self.assertIn("a.csv: empty", self.messages(d))
        with open(os.path.join(d, "a.csv"), "w", encoding="utf-8") as fh:
            fh.write("x,\n1,2\n")
        self.assertIn("a header cell is empty", self.messages(d))

    def test_file_column_must_be_a_plain_csv_name(self):
        d = self.good()
        with open(os.path.join(d, "index.csv"), encoding="utf-8") as fh:
            text = fh.read().replace("a.csv", "../a.csv", 1)
        with open(os.path.join(d, "index.csv"), "w", encoding="utf-8") as fh:
            fh.write(text)
        self.assertIn("must be a plain file name", self.messages(d))


class BuildUpdatedRule(Base):
    """`updated` = the date the CSV's cells last changed, as recorded by build (docs/DECISIONS.md)."""

    def test_updated_moves_only_when_content_changes(self):
        d = make_set(os.path.join(self.tmp, "u"), {"a.csv": (["x"], [["1"]]), "b.csv": (["y"], [["2"]])})
        self.build(d, "2026-10-08")
        self.assertEqual(set(r["updated"] for r in sh.read_index(d)[1]), {"2026-10-08"})
        self.build(d, "2026-10-20")                       # nothing changed: dates stay
        self.assertEqual(set(r["updated"] for r in sh.read_index(d)[1]), {"2026-10-08"})
        with open(os.path.join(d, "a.csv"), "a", encoding="utf-8") as fh:
            fh.write("9\n")
        self.build(d, "2026-10-21")
        got = dict((r["file"], (r["updated"], r["rows"])) for r in sh.read_index(d)[1])
        self.assertEqual(got, {"a.csv": ("2026-10-21", "2"), "b.csv": ("2026-10-08", "1")})
        self.assertEqual(sh.lint_set(d), [])

    def test_line_endings_and_quoting_are_not_a_change(self):
        d = make_set(os.path.join(self.tmp, "crlf"), {"a.csv": (["x", "y"], [["1", "a,b"]])})
        self.build(d, "2026-10-08")
        with open(os.path.join(d, "a.csv"), "w", encoding="utf-8", newline="") as fh:
            fh.write('"x","y"\r\n"1","a,b"\r\n')
        self.build(d, "2026-11-01")
        self.assertEqual(sh.read_index(d)[1][0]["updated"], "2026-10-08")

    def test_first_build_keeps_a_valid_hand_set_date(self):
        d = make_set(os.path.join(self.tmp, "hand"), {"a.csv": (["x"], [["1"]])})
        rows = sh.read_index(d)[1]
        rows[0]["updated"] = "2026-01-02"
        with open(os.path.join(d, "index.csv"), "w", encoding="utf-8", newline="") as fh:
            fh.write(sh.write_csv_text(sh.INDEX_COLS, [[r[c] for c in sh.INDEX_COLS] for r in rows]))
        self.build(d, "2026-10-08")
        self.assertEqual(sh.read_index(d)[1][0]["updated"], "2026-01-02")

    def test_build_refuses_a_broken_set(self):
        d = make_set(os.path.join(self.tmp, "broken"), {"a.csv": (["x"], [])})
        os.remove(os.path.join(d, "a.csv"))
        problems, plan = sh.plan_build(d, "2026-10-08")
        self.assertIsNone(plan)
        self.assertTrue(any("does not exist" in p for p in problems))


class GeneratedMachSet(unittest.TestCase):
    """The committed docs/sheets set and the generator that writes it."""

    @classmethod
    def setUpClass(cls):
        cls.out, cls.md = gen.generate(ROOT, None)

    def test_committed_set_lints_clean(self):
        self.assertEqual(sh.lint_set(SET_DIR), [])

    def test_every_index_row_is_generated(self):
        self.assertEqual([m["file"] for m in gen.index_meta()], list(self.out))

    def test_commands_cover_every_libexec_command(self):
        rows = self.out["commands.csv"][1]
        have = set(r[0] for r in rows)
        for fn in os.listdir(os.path.join(ROOT, "libexec")):
            if fn.startswith("mach-") and "." not in fn:
                self.assertIn(fn[5:], have, fn)
        self.assertIn("<subsystem>", have)

    def test_declared_subcommands_all_have_a_usage_line(self):
        for r in self.out["commands.csv"][1]:
            self.assertTrue(r[2], "no usage for %s %s" % (r[0], r[1]))
            self.assertIn(r[5], gen.STATUSES)
            self.assertIn(r[4], ("upkeep", "gateway", "subsystem"))
            self.assertIn(r[6], ("yes", "no"))

    def test_status_and_apply_columns(self):
        rows = dict(((r[0], r[1]), r) for r in self.out["commands.csv"][1])
        self.assertEqual(rows[("memory", "collect")][5], "preview-only")
        self.assertEqual(rows[("mcp", "apply")][5], "stub")
        self.assertEqual(rows[("plans", "archive")][7], "yes")
        self.assertEqual(rows[("plans", "list")][7], "no")
        self.assertEqual(rows[("plans", "list")][6], "no")
        self.assertEqual(rows[("subsystem", "unit")][7], "yes")      # tail shared by `unit|role <id> --from D [--apply]`
        self.assertEqual(rows[("sheets", "build")][8], "yes")

    def test_every_flag_has_a_meaning(self):
        missing = [(r[0], r[1]) for r in self.out["flags.csv"][1] if not r[3]]
        self.assertEqual(missing, [], "add a `# mach-flag: --flag | meaning | default` header to the command (or COMMON_FLAGS)")

    def test_flag_arguments_are_found(self):
        flags = dict(((r[0], r[1]), r) for r in self.out["flags.csv"][1])
        self.assertEqual(flags[("do", "--tool")][2], "T")
        self.assertEqual(flags[("research", "--depth")][2], "quick|deep")
        self.assertEqual(flags[("(global)", "--host")][2], "H")
        self.assertIn(("sheets", "--check"), flags)

    def test_functions_cover_python_and_bash(self):
        langs = set(r[2] for r in self.out["functions.csv"][1])
        self.assertEqual(langs, {"python", "bash"})
        have = set((r[0], r[1]) for r in self.out["functions.csv"][1])
        self.assertTrue(("lib/mach_sheets.py", "build_xlsx") in have)
        self.assertTrue(("lib/common.sh", "err") in have)
        self.assertTrue(all(r[0].startswith(("lib/", "libexec/")) for r in self.out["functions.csv"][1]))

    def test_config_keys_come_with_defaults_and_meanings(self):
        rows = dict(((r[0], r[1]), r) for r in self.out["config-keys.csv"][1])
        self.assertEqual(rows[("doctor", "disk_min_free_pct")][2], "10")
        self.assertTrue(rows[("note", "default")][3])
        self.assertIn(("env", "MACH_NOTE_DEFAULT"), rows)

    def test_destinations_name_every_tool(self):
        have = set(r[0] for r in self.out["destinations.csv"][1])
        for t in ("claude", "codex", "cursor-agent", "grokcli-agent", "kiro", "user", "grokbot:<Service>", "mach note inbox"):
            self.assertIn(t, have)

    def test_sources_are_generic_and_in_repo(self):
        units = dict((r[0], r) for r in self.out["subsystem-units.csv"][1])
        self.assertEqual(sorted(units), ["alpha", "beta", "gamma"])
        self.assertEqual(units["beta"][2], "alpha")
        names = [r[0] for r in self.out["templates.csv"][1]]
        self.assertIn("~/machine-flows/ai-authored/templates/task-brief.md", names)
        self.assertTrue(all(n.startswith("~/machine-flows/ai-authored/templates/") for n in names))

    def test_no_absolute_home_paths_in_generated_files_or_workbook(self):
        bad = re.compile(r"/home/|/Users/")
        for fn in os.listdir(SET_DIR):
            p = os.path.join(SET_DIR, fn)
            if fn.endswith(".xlsx"):
                with zipfile.ZipFile(p) as z:
                    text = b"".join(z.read(n) for n in z.namelist()).decode("utf-8")
            else:
                with open(p, encoding="utf-8") as fh:
                    text = fh.read()
            self.assertIsNone(bad.search(text), fn)
        self.assertIsNone(bad.search(self.md))
        for header, rows in [v for v in self.out.values() if v]:
            for r in rows:
                for cell in r:
                    self.assertIsNone(bad.search(cell), cell)

    def test_markdown_table_has_both_tables(self):
        self.assertIn("## Commands", self.md)
        self.assertIn("## Flags", self.md)
        self.assertIn("| command | subcommand | usage |", self.md)
        self.assertIn("quick\\|deep", self.md)       # pipes inside cells are escaped


class UsageParsing(unittest.TestCase):
    def test_split_top_respects_brackets(self):
        self.assertEqual(gen.split_top("a|b <x|y>|c [d|e]", "|"), ["a", "b <x|y>", "c [d|e]"])

    def test_flag_regex(self):
        found = dict((m.group(1), m.group(2) or m.group(3) or "") for m in gen.FLAG_RE.finditer(
            "[--tool T] [--depth quick|deep] [--light|--full] [--role R|--to D] [-t TITLE] [--tag X ...] --from <file>"))
        self.assertEqual(found, {"--tool": "T", "--depth": "quick|deep", "--light": "", "--full": "", "--role": "R", "--to": "D",
                                 "-t": "TITLE", "--tag": "X", "--from": "<file>"})

    def test_header_continuation_lines_join(self):
        d = scratch()
        try:
            p = os.path.join(d, "mach-x")
            with open(p, "w") as fh:
                fh.write('#!/usr/bin/env python3\n# mach-summary: s\n# mach-usage: mach x "a" [-t T]\n# mach-usage:        [--json]\n# mach-usage: mach x y\n')
            h = gen.headers(p)
            self.assertEqual(h["usage"], ['mach x "a" [-t T] [--json]', "mach x y"])
        finally:
            shutil.rmtree(d, True)


if __name__ == "__main__":
    unittest.main()
